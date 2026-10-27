"""Follow real site structure into searchable evidence without a search API."""
from __future__ import annotations

import asyncio
from io import BytesIO
import json

import httpx
import pytest


def pdf_bytes(text="robot admission evidence", pages=1):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for number in range(pages):
        page = writer.add_blank_page(width=595, height=842)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
            DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 12 Tf 20 800 Td (" + text + " page " + str(number + 1) + ") Tj ET").encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def xlsx_bytes():
    from openpyxl import Workbook
    workbook = Workbook()
    first = workbook.active
    first.title = "2026 admissions"
    first.append(["program", "score", "quota"])
    first.append(["robot computing", 365, 12, "=1+2"])
    second = workbook.create_sheet("2025 admissions")
    second.append(["program", "score"])
    second.append(["robot computing", 351])
    out = BytesIO()
    workbook.save(out)
    workbook.close()
    return out.getvalue()


def test_xlsx_extracts_sheets_rows_and_cell_coordinates():
    from periscope.collection.documents import extract_document
    doc = extract_document(xlsx_bytes(), "application/octet-stream",
                           "https://lab.example/admissions.xlsx")
    assert doc.format == "xlsx"
    assert "robot computing" in doc.text
    row = next(row for row in doc.metadata()["sheet_rows"]
               if row["sheet"] == "2026 admissions" and row["row"] == 2)
    line = doc.text[row["start"]:row["end"]]
    assert "A2=robot computing" in line
    assert "B2=365" in line and "C2=12" in line
    assert "D2==1+2" in line
    assert any(row["sheet"] == "2025 admissions" for row in doc.metadata()["sheet_rows"])


def test_xlsx_text_limit_is_reported():
    from periscope.collection.documents import extract_document
    doc = extract_document(xlsx_bytes(), "application/octet-stream",
                           "https://lab.example/admissions.xlsx", max_chars=20)
    assert len(doc.text) <= 20
    assert doc.truncated


def test_xlsx_rejects_archive_with_excessive_entries():
    from zipfile import ZipFile
    from periscope.collection.documents import extract_document
    out = BytesIO()
    with ZipFile(out, "w") as archive:
        for index in range(4097):
            archive.writestr("empty/%d" % index, "")
    with pytest.raises(ValueError, match="xlsx_expanded_size_limit"):
        extract_document(out.getvalue(), "application/octet-stream",
                         "https://lab.example/admissions.xlsx")


def test_xlsx_empty_sheet_without_dimension_does_not_hide_next_sheet():
    from openpyxl import Workbook
    from zipfile import ZipFile
    import re
    from periscope.collection.documents import extract_document
    workbook = Workbook()
    workbook.create_sheet("Admissions").append(["robot computing", 365])
    original = BytesIO()
    workbook.save(original)
    workbook.close()
    rewritten = BytesIO()
    with ZipFile(BytesIO(original.getvalue())) as source, ZipFile(rewritten, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            target.writestr(entry, data)
    doc = extract_document(rewritten.getvalue(), "application/octet-stream",
                           "https://lab.example/admissions.xlsx")
    assert "robot computing" in doc.text
    assert not doc.truncated


def test_html_extraction_keeps_body_and_resolves_child_links():
    from periscope.collection.documents import extract_document
    html = '<html><title>Robot admissions</title><nav>noise</nav><main><h1>Robot admissions</h1><p>' + "Independent research evidence. " * 70 + '</p><a href="../files/admissions.pdf">Admission attachment</a></main></html>'
    doc = extract_document(html.encode(), "text/html", "https://lab.example/news/1.html")
    assert "Independent research evidence" in doc.text
    assert doc.format == "html"
    assert any(link.url == "https://lab.example/files/admissions.pdf" for link in doc.links)
    assert "noise" not in doc.text


def test_pdf_extraction_has_page_locators_and_honest_limits():
    from periscope.collection.documents import extract_document
    doc = extract_document(pdf_bytes(pages=2), "application/pdf", "https://lab.example/record.pdf")
    assert "robot admission evidence" in doc.text
    assert [span["page"] for span in doc.page_spans] == [1, 2]
    assert all(doc.text[span["start"]:span["end"]] for span in doc.page_spans)
    limited = extract_document(pdf_bytes(pages=2), "application/pdf", "https://lab.example/record.pdf", max_chars=12)
    assert len(limited.text) == 12 and limited.truncated


def test_image_only_pdf_is_marked_as_needing_ocr():
    from periscope.collection.documents import extract_document
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    out = BytesIO()
    writer.write(out)
    doc = extract_document(out.getvalue(), "application/pdf", "https://lab.example/scan.pdf")
    assert not doc.text and doc.needs_ocr


def test_sitemap_is_discovery_instead_of_fake_article_text():
    from periscope.collection.documents import extract_document
    doc = extract_document(b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://lab.example/post</loc></url></urlset>',
                           "application/xml", "https://lab.example/sitemap.xml")
    assert doc.format == "sitemap" and not doc.text
    assert doc.links[0].url == "https://lab.example/post"


async def crawl_fixture(monkeypatch, responses, **options):
    from periscope.collection import crawler
    called = []
    async def public(url):
        if "127.0.0.1" in url:
            from src.url_security import UnsafeURLError
            raise UnsafeURLError("private address")
        return url
    monkeypatch.setattr(crawler, "validate_public_http_url", public)
    async def handler(request):
        called.append(str(request.url))
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text=options.get("robots", "User-agent: *\nAllow: /"))
        value = responses.get(request.url.path, httpx.Response(404))
        if callable(value):
            return await value(request)
        return value
    settings = crawler.CrawlOptions(seeds=tuple(options.pop("seeds", ("https://lab.example/list",))),
                                   max_pages=options.pop("max_pages", 8), depth=options.pop("depth", 2),
                                   budget=options.pop("budget", 8), delay=0, max_bytes=options.pop("max_bytes", 1024 * 1024))
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        rows, stat = await crawler.crawl_async("robot", settings, client=client)
    return rows, stat, called


def test_listing_discovers_article_and_pdf_with_parent_evidence(monkeypatch):
    listing = '<main><a href="/article">Robot research report</a><a href="/table.pdf">Robot table</a><a href="https://other.example/offsite">external</a></main>'
    article = '<main><h1>Robot report</h1><p>' + "robot research findings. " * 60 + '</p></main>'
    rows, stat, called = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, text=listing, headers={"content-type": "text/html"}),
        "/article": httpx.Response(200, text=article, headers={"content-type": "text/html"}),
        "/table.pdf": httpx.Response(200, content=pdf_bytes(), headers={"content-type": "application/pdf"})}))
    by_url = {row["url"]: row for row in rows}
    assert {"https://lab.example/article", "https://lab.example/table.pdf"} <= by_url.keys()
    pdf = by_url["https://lab.example/table.pdf"]
    assert pdf["provenance"]["parent_url"] == "https://lab.example/list"
    assert pdf["document"]["page_spans"][0]["page"] == 1
    assert all(row["provenance"]["identity"] == "url" for row in rows)
    assert not any("other.example" in url for url in called)
    assert stat["rows"] == len(rows)


def test_listing_discovers_xlsx_with_source_and_row_locations(monkeypatch, scratch_path):
    listing = '<main><a href="/admissions.xlsx">robot admission spreadsheet</a></main>'
    rows, stat, _ = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, text=listing, headers={"content-type": "text/html"}),
        "/admissions.xlsx": httpx.Response(200, content=xlsx_bytes(),
                                           headers={"content-type": "application/octet-stream"})}))
    sheet = next(row for row in rows if row["url"].endswith("admissions.xlsx"))
    assert sheet["provenance"]["parent_url"] == "https://lab.example/list"
    assert sheet["document"]["format"] == "xlsx"
    assert any(item["row"] == 2 for item in sheet["document"]["sheet_rows"])
    assert any(report["status"] == "ok" and report["format"] == "xlsx"
               for report in stat["reports"])
    from periscope.operations.ingest import build
    from periscope.topics import runs
    from src.corpus.store import Corpus
    path = scratch_path / "xlsx.db"
    items, skipped = build(runs.ingestable_rows({"rows": [runs._row_for_disk(sheet, "crawl")]}), [])
    assert not skipped
    corpus = Corpus(path)
    corpus.add_items(items)
    corpus.close()
    corpus = Corpus(path)
    try:
        assert corpus.search("robot computing")
        saved = corpus._conn.execute("SELECT metadata_json FROM items WHERE url=?", (sheet["url"],)).fetchone()
        metadata = json.loads(saved["metadata_json"])
        assert any(row["sheet"] == "2026 admissions" and row["row"] == 2
                   for row in metadata["document"]["sheet_rows"])
    finally:
        corpus.close()


def test_truncated_xlsx_is_reported_as_partial(monkeypatch):
    from openpyxl import Workbook
    from periscope.topics.models import TopicRequest
    from periscope.topics.service import TopicService
    workbook = Workbook()
    workbook.active.append(["robot " + "x" * 32000, "y" * 32000])
    out = BytesIO()
    workbook.save(out)
    workbook.close()
    rows, stat, _ = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, content=out.getvalue(), headers={
            "content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})}))
    assert rows and rows[0]["document"]["text_truncated"]
    assert stat["reports"][0]["status"] == "partial"
    service = TopicService(crawl_lookup=lambda *a, **kw: (rows, stat))
    result = service.search(TopicRequest(topic="robot", no_local=True, no_api=True, speed="off",
                                         crawl_seeds=("https://lab.example/list",)))
    assert result.source_reports[0]["status"] == "partial"
    from periscope.topics.render import render
    display = render("robot", rows, {})
    assert "Excel" in display and "已截断" in display


def test_robots_denial_is_visible_and_never_fetched(monkeypatch):
    rows, stat, called = asyncio.run(crawl_fixture(monkeypatch, {}, robots="User-agent: *\nDisallow: /list"))
    assert not rows
    assert "https://lab.example/list" not in called
    assert stat["reports"][0]["status"] == "blocked"


def test_redirect_cannot_escape_seed_scope(monkeypatch):
    _, stat, called = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(302, headers={"location": "https://other.example/private"})}))
    assert not any("other.example" in url for url in called)
    assert stat["reports"][0]["status"] == "blocked"


def test_failed_page_does_not_erase_completed_evidence(monkeypatch):
    html = '<main><p>' + "robot research. " * 70 + '</p><a href="/broken">Robot next report</a></main>'
    rows, stat, _ = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, text=html, headers={"content-type": "text/html"}),
        "/broken": httpx.Response(503)}))
    assert rows
    assert any(report["status"] == "failed" for report in stat["reports"])


def test_budget_retains_finished_pages_and_reports_unfinished(monkeypatch):
    async def slow(request):
        await asyncio.sleep(5)
        return httpx.Response(200)
    html = '<main><p>' + "robot research. " * 70 + '</p><a href="/slow">Robot second report</a></main>'
    rows, stat, _ = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, text=html, headers={"content-type": "text/html"}), "/slow": slow}, budget=2))
    assert rows
    assert any(report["status"] == "over_budget" for report in stat["reports"])


def test_page_limit_counts_nonmatching_pages_too(monkeypatch):
    html = '<main><a href="/second">robot research</a></main>'
    _, stat, called = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, text=html, headers={"content-type": "text/html"})}, max_pages=1))
    assert "https://lab.example/second" not in called
    assert stat["pages"] == 1


def test_oversized_document_is_reported_without_indexing(monkeypatch):
    rows, stat, _ = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, content=b"robot " * 1000, headers={"content-type": "text/html"})}, max_bytes=100))
    assert not rows
    assert stat["reports"][0]["status"] == "too_large"


def test_private_seed_is_rejected_before_any_request(monkeypatch):
    rows, stat, called = asyncio.run(crawl_fixture(monkeypatch, {}, seeds=("http://127.0.0.1/list",)))
    assert not rows and not called
    assert stat["reports"][0]["status"] == "blocked"


def test_topic_service_can_crawl_with_search_api_disabled():
    from periscope.topics.models import TopicRequest
    from periscope.topics.service import TopicService
    row = {"title": "robot evidence", "url": "https://lab.example/report", "body": "robot " * 100,
           "text": "robot " * 100, "source": "web_crawl", "kind": "media", "_origin": "crawl"}
    service = TopicService(crawl_lookup=lambda *a, **kw: ([row], {"reports": [{"status": "ok"}], "elapsed": 0.1}))
    result = service.search(TopicRequest(topic="robot", no_local=True, no_api=True, speed="off",
                                         crawl_seeds=("https://lab.example/list",)))
    assert result.crawl_rows == [row]
    assert result.merged[0]["url"] == row["url"]
    assert result.coverage["direct_linked_urls"] == 1
    assert result.coverage["direct_urls_not_in_api"] is None


def test_crawl_metadata_and_changed_versions_survive_save_and_retrieval(scratch_path, monkeypatch):
    from periscope.operations.ingest import build
    from periscope.topics import layers, runs
    from src.corpus.store import Corpus
    base = {"title": "robot admission record", "url": "https://lab.example/record.pdf", "source": "web_crawl",
            "kind": "media", "body": "robot research evidence " * 50, "_origin": "crawl",
            "provenance": {"identity": "url", "method": "crawl", "parent_url": "https://lab.example/list"}}
    corpus = Corpus(scratch_path / "corpus.db")
    for fingerprint in ("old", "new"):
        document = {"format": "pdf", "sha256": fingerprint, "page_spans": [{"page": 1, "start": 0, "end": 100}]}
        payload = {"rows": [runs._row_for_disk({**base, "document": document}, "crawl")]}
        items, skipped = build(runs.ingestable_rows(payload), [])
        assert not skipped
        assert items[0].metadata["document"] == document
        corpus.add_items(items)
    assert corpus._conn.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 2
    monkeypatch.setattr(layers.corpus_search, "CORPUS", scratch_path / "corpus.db")
    rows, _ = layers.local_lookup("robot", 5)
    assert rows[0]["document"]["format"] == "pdf"
    corpus.close()


def test_redirect_also_obeys_robots_for_its_destination(monkeypatch):
    rows, stat, called = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(302, headers={"location": "/private"}),
        "/private": httpx.Response(200, text="robot private")}, robots="User-agent: *\nDisallow: /private"))
    assert not rows
    assert "https://lab.example/private" not in called
    assert stat["reports"][0]["status"] == "blocked"


def test_crawler_failure_is_not_summarized_as_no_matches():
    from periscope.topics.models import TopicRequest
    from periscope.topics.service import TopicService
    service = TopicService(crawl_lookup=lambda *a, **kw: ([], {"elapsed": 0.1,
                         "reports": [{"status": "failed", "error": "HTTP_503"}]}))
    result = service.search(TopicRequest(topic="robot", no_local=True, no_api=True, speed="off",
                                         crawl_seeds=("https://lab.example/list",)))
    assert result.source_reports[0]["status"] == "failed"


def test_incremental_ingest_keeps_versions_in_one_run(scratch_path, monkeypatch):
    from periscope.operations import ingest
    from src.corpus.store import Corpus
    rows = [{"title": "robot publication", "url": "https://lab.example/report", "source": "web_crawl",
             "kind": "media", "text": "robot findings " * 60, "document": {"sha256": version}}
            for version in ("first", "second")]
    file = scratch_path / "http_items.crawler.json"
    file.write_text(json.dumps(rows), encoding="utf-8")
    monkeypatch.setattr(ingest, "CORPUS", scratch_path / "versions.db")
    monkeypatch.setattr(ingest, "fulltext_cache", lambda: scratch_path / "no-cache.jsonl")
    monkeypatch.setattr(ingest.sys, "argv", ["ingest", "--only", str(file)])
    assert ingest.main() == 0
    corpus = Corpus(scratch_path / "versions.db")
    assert corpus._conn.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 2
    corpus.close()


def test_pdf_worker_returns_page_evidence():
    from periscope.collection.documents import extract_document_async
    doc = asyncio.run(extract_document_async(pdf_bytes(), "application/pdf", "https://lab.example/report.pdf"))
    assert doc.page_spans[0]["page"] == 1
    assert "robot admission evidence" in doc.text


def test_cancelled_document_extraction_stops_its_worker(monkeypatch):
    from periscope.collection.documents import extract_document_async
    class Worker:
        returncode = None
        killed = False
        async def communicate(self, data):
            await asyncio.sleep(5)
        def kill(self):
            self.killed = True
            self.returncode = -9
        async def wait(self):
            return self.returncode
    worker = Worker()
    async def spawn(*args, **kwargs):
        return worker
    monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
    async def exercise():
        with pytest.raises(asyncio.TimeoutError):
            await asyncio.wait_for(extract_document_async(b"data", "text/html", "https://lab.example/"), 0.03)
    asyncio.run(exercise())
    assert worker.killed


def test_auto_seed_selection_spreads_work_across_hosts():
    from periscope.collection.crawler import select_seeds
    rows = [{"url": "https://one.example/" + str(i)} for i in range(20)] + [{"url": "https://two.example/research"}]
    seeds = select_seeds(rows, limit=2)
    assert seeds == ("https://one.example/0", "https://two.example/research")


def test_auto_crawl_follows_direct_results_without_the_api():
    from periscope.topics.models import TopicRequest
    from periscope.topics.service import TopicService
    seen = []
    def crawl(topic, seeds, **kwargs):
        seen.extend(seeds)
        return [], {"elapsed": 0.1, "reports": []}
    platform = lambda *a, **kw: ([{"title": "robot research", "url": "https://lab.example/article", "source": "institution_feed"}],
                                {"failed": [], "elapsed": 0.1})
    service = TopicService(platform_lookup=platform, crawl_lookup=crawl)
    service.search(TopicRequest(topic="robot", no_local=True, no_api=True, crawl=True))
    assert seen == ["https://lab.example/article"]


def test_robots_declared_sitemap_discovers_otherwise_unlinked_pages(monkeypatch):
    rows, stat, _ = asyncio.run(crawl_fixture(monkeypatch, {
        "/list": httpx.Response(200, text="<main>No matching articles on this page.</main>", headers={"content-type": "text/html"}),
        "/sitemap.xml": httpx.Response(200, text='<urlset><url><loc>https://lab.example/hidden-report</loc></url></urlset>', headers={"content-type": "application/xml"}),
        "/hidden-report": httpx.Response(200, text='<main><p>' + 'robot original study. '*50 + '</p></main>', headers={"content-type": "text/html"})
    }, robots="User-agent: *\nAllow: /\nSitemap: https://lab.example/sitemap.xml"))
    assert any(row["url"] == "https://lab.example/hidden-report" for row in rows)
